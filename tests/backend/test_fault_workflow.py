"""Fault repair workflow: assignment, due dates, priority, status transitions,
resolved_at bookkeeping, repair proof photos and work-order export."""
from __future__ import annotations

import csv
import io

import pytest

from ml.src.utils import SEVERITY_MAP

# Smallest valid file signatures — the upload endpoint sniffs magic bytes.
PNG_BYTES = b"\x89PNG\r\n\x1a\n" + b"\x00" * 64
JPEG_BYTES = b"\xff\xd8\xff\xe0" + b"\x00" * 64


def _seed_fault(session, park_id="PARK_WF", panel_id="R1-C1", cls="bypass-diode",
                status="open", gps='{"lat": 18.5, "lon": 73.8}', **extra) -> int:
    from axalon.db.models import PanelFault, Park

    if session.query(Park).filter_by(id=park_id).first() is None:
        session.add(Park(id=park_id, name="Workflow Park"))
        session.flush()
    fault = PanelFault(
        park_id=park_id,
        panel_id=panel_id,
        class_=cls,
        class_id=4,
        severity=SEVERITY_MAP.get(cls),
        status=status,
        last_gps=gps,
        **extra,
    )
    session.add(fault)
    session.commit()
    return fault.id


@pytest.fixture
def fault_id(db_session) -> int:
    return _seed_fault(db_session)


@pytest.fixture
def local_store(monkeypatch, tmp_path):
    """Force the local-disk object-store fallback into a temp directory."""
    for var in ("AZURE_STORAGE_CONNECTION_STRING", "SUPABASE_URL", "SUPABASE_SERVICE_KEY"):
        monkeypatch.delenv(var, raising=False)
    from axalon.api.routers import faults as faults_router

    photos_dir = tmp_path / "fault_photos"  # tmp_path also holds the test DB
    photos_dir.mkdir()
    monkeypatch.setattr(faults_router, "FAULT_PHOTOS_DIR", photos_dir)
    return photos_dir


# ── serializer fields & priority ──────────────────────────────────────────────

def test_fault_exposes_workflow_fields(client, fault_id):
    fault = client.get("/parks/PARK_WF/faults").json()["faults"][0]
    for key in ("assignee", "due_date", "priority", "priority_override",
                "resolved_at", "resolution_note", "photo_count"):
        assert key in fault
    assert fault["assignee"] is None
    assert fault["resolved_at"] is None


def test_priority_is_derived_from_class_severity_when_not_set(client, db_session):
    _seed_fault(db_session, panel_id="R2-C2", cls="soiling")
    _seed_fault(db_session, panel_id="R3-C3", cls="hot-spot-high")
    faults = {f["class"]: f for f in client.get("/parks/PARK_WF/faults").json()["faults"]}
    assert SEVERITY_MAP["soiling"] == "LOW"
    assert faults["soiling"]["priority"] == "low"
    assert faults["hot-spot-high"]["priority"] == "urgent"
    assert faults["soiling"]["priority_override"] is None


def test_priority_override_wins_and_can_be_cleared(client, fault_id):
    r = client.patch(f"/faults/{fault_id}", json={"priority": "low"})
    assert r.status_code == 200, r.text
    assert r.json()["priority"] == "low"
    assert r.json()["priority_override"] == "low"

    r = client.patch(f"/faults/{fault_id}", json={"priority": None})
    assert r.status_code == 200
    assert r.json()["priority_override"] is None
    assert r.json()["priority"] == "urgent"  # bypass-diode → CRITICAL


def test_invalid_priority_rejected(client, fault_id):
    r = client.patch(f"/faults/{fault_id}", json={"priority": "whenever"})
    assert r.status_code == 400
    assert "priority" in r.json()["detail"].lower()


# ── assignment & due date ─────────────────────────────────────────────────────

def test_assign_with_due_date(client, fault_id):
    r = client.patch(f"/faults/{fault_id}", json={
        "status": "assigned", "assignee": "ravi@om.example", "due_date": "2026-10-01",
    })
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["status"] == "assigned"
    assert body["assignee"] == "ravi@om.example"
    assert body["due_date"] == "2026-10-01"


def test_bad_due_date_rejected(client, fault_id):
    r = client.patch(f"/faults/{fault_id}", json={"due_date": "next tuesday"})
    assert r.status_code == 400
    assert "due_date" in r.json()["detail"]


def test_assigned_status_requires_an_assignee(client, fault_id):
    r = client.patch(f"/faults/{fault_id}", json={"status": "assigned"})
    assert r.status_code == 422
    assert "assignee" in r.json()["detail"].lower()


def test_cannot_clear_assignee_while_work_is_in_flight(client, fault_id):
    client.patch(f"/faults/{fault_id}", json={"status": "assigned", "assignee": "ravi"})
    r = client.patch(f"/faults/{fault_id}", json={"assignee": None})
    assert r.status_code == 422


# ── status transitions ────────────────────────────────────────────────────────

@pytest.mark.parametrize("path", [
    ["assigned", "in_progress", "resolved"],
    ["in_progress", "resolved", "open"],
    ["resolved"],
    ["stale", "open"],
    ["assigned", "open"],
])
def test_valid_transitions(client, fault_id, path):
    client.patch(f"/faults/{fault_id}", json={"assignee": "crew-a"})
    for status in path:
        r = client.patch(f"/faults/{fault_id}", json={"status": status})
        assert r.status_code == 200, f"{status}: {r.text}"
        assert r.json()["status"] == status


@pytest.mark.parametrize("start,target", [
    ("resolved", "in_progress"),
    ("resolved", "assigned"),
    ("resolved", "stale"),
    ("in_progress", "stale"),
    ("assigned", "stale"),
])
def test_invalid_transitions_return_422(client, db_session, start, target):
    fid = _seed_fault(db_session, panel_id="R7-C7", status=start, assignee="crew-a")
    r = client.patch(f"/faults/{fid}", json={"status": target})
    assert r.status_code == 422, r.text
    detail = r.json()["detail"]
    assert start in detail and target in detail


def test_unknown_status_returns_400(client, fault_id):
    r = client.patch(f"/faults/{fault_id}", json={"status": "done-ish"})
    assert r.status_code == 400


def test_same_status_is_a_noop(client, fault_id):
    r = client.patch(f"/faults/{fault_id}", json={"status": "open"})
    assert r.status_code == 200


def test_legacy_statuses_still_listable(client, db_session):
    _seed_fault(db_session, panel_id="R4-C4", status="stale")
    _seed_fault(db_session, panel_id="R5-C5", status="resolved")
    body = client.get("/parks/PARK_WF/faults").json()
    assert body["counts_by_status"]["stale"] == 1
    assert body["counts_by_status"]["resolved"] == 1
    assert body["counts_by_status"]["in_progress"] == 0
    assert client.get("/parks/PARK_WF/faults?status=in_progress").status_code == 200


# ── resolved_at ───────────────────────────────────────────────────────────────

def test_resolved_at_set_on_resolve_and_cleared_on_reopen(client, fault_id):
    r = client.patch(f"/faults/{fault_id}", json={
        "status": "resolved", "resolution_note": "Replaced bypass diode",
    })
    assert r.status_code == 200
    assert r.json()["resolved_at"] is not None
    assert r.json()["resolution_note"] == "Replaced bypass diode"
    first_resolved_at = r.json()["resolved_at"]

    # re-saving a resolved fault does not move the timestamp
    r = client.patch(f"/faults/{fault_id}", json={"resolution_note": "Replaced diode + cleaned"})
    assert r.json()["resolved_at"] == first_resolved_at

    r = client.patch(f"/faults/{fault_id}", json={"status": "open"})
    assert r.status_code == 200
    assert r.json()["resolved_at"] is None


def test_redetection_keeps_in_flight_work_but_reopens_resolved(db_session):
    from axalon.db.models import Inspection, PanelFault
    from axalon.pipeline.tracking import reconcile_inspection

    busy = _seed_fault(db_session, panel_id="R1-C1", cls="bypass-diode",
                       status="in_progress", assignee="crew-a")
    from datetime import datetime
    done = _seed_fault(db_session, panel_id="R2-C2", cls="soiling",
                       status="resolved", resolved_at=datetime(2026, 9, 1))
    db_session.add(Inspection(id="insp-wf", park_id="PARK_WF"))
    db_session.commit()
    dets = [
        {"panel_id": "R1-C1", "class": "bypass-diode", "severity": "CRITICAL", "confidence": 0.9},
        {"panel_id": "R2-C2", "class": "soiling", "severity": "LOW", "confidence": 0.9},
    ]
    reconcile_inspection(db_session, "PARK_WF", "insp-wf", "2026-09-10", dets)

    busy_f = db_session.query(PanelFault).filter_by(id=busy).one()
    done_f = db_session.query(PanelFault).filter_by(id=done).one()
    assert busy_f.status == "in_progress"
    assert done_f.status == "open"
    assert done_f.resolved_at is None


# ── repair proof photos ───────────────────────────────────────────────────────

def test_upload_list_and_serve_photo(client, fault_id, local_store):
    r = client.post(
        f"/faults/{fault_id}/photos",
        files={"file": ("../../etc/after repair.png", PNG_BYTES, "image/png")},
    )
    assert r.status_code == 201, r.text
    photo = r.json()
    assert photo["fault_id"] == fault_id
    assert photo["original_name"] == "after repair.png"
    assert photo["content_type"] == "image/png"
    assert photo["size_bytes"] == len(PNG_BYTES)

    listed = client.get(f"/faults/{fault_id}/photos").json()
    assert [p["id"] for p in listed] == [photo["id"]]

    served = client.get(f"/faults/{fault_id}/photos/{photo['id']}")
    assert served.status_code == 200
    assert served.content == PNG_BYTES
    assert served.headers["content-type"] == "image/png"
    assert served.headers.get("x-content-type-options") == "nosniff"

    assert client.get("/parks/PARK_WF/faults").json()["faults"][0]["photo_count"] == 1


def test_photo_filename_is_header_safe(client, fault_id, local_store):
    r = client.post(
        f"/faults/{fault_id}/photos",
        files={"file": ('x";evil=1.jpg', JPEG_BYTES, "image/jpeg")},
    )
    assert r.status_code == 201, r.text
    assert '"' not in r.json()["original_name"] and ";" not in r.json()["original_name"]
    served = client.get(f"/faults/{fault_id}/photos/{r.json()['id']}")
    assert served.headers["content-disposition"].count('"') == 2


def test_photo_rejects_non_image_content_type(client, fault_id, local_store):
    r = client.post(
        f"/faults/{fault_id}/photos",
        files={"file": ("notes.pdf", b"%PDF-1.4 hello", "application/pdf")},
    )
    assert r.status_code == 415


def test_photo_rejects_spoofed_image(client, fault_id, local_store):
    r = client.post(
        f"/faults/{fault_id}/photos",
        files={"file": ("evil.png", b"<script>alert(1)</script>", "image/png")},
    )
    assert r.status_code == 415


def test_photo_rejects_oversize(client, fault_id, local_store, monkeypatch):
    from axalon.api.routers import faults as faults_router

    monkeypatch.setattr(faults_router, "_MAX_FAULT_PHOTO_BYTES", 100)
    r = client.post(
        f"/faults/{fault_id}/photos",
        files={"file": ("big.jpg", JPEG_BYTES + b"\x00" * 200, "image/jpeg")},
    )
    assert r.status_code == 413
    assert list(local_store.iterdir()) == []  # nothing left behind


def test_photo_upload_unknown_fault_404(client, local_store):
    r = client.post("/faults/9999/photos", files={"file": ("a.png", PNG_BYTES, "image/png")})
    assert r.status_code == 404


def test_photo_uses_object_store_when_configured(client, fault_id, local_store, monkeypatch):
    from axalon.api.routers import faults as faults_router

    stored: dict[str, bytes] = {}

    class FakeStore:
        def upload(self, name, fileobj, content_type):
            stored[name] = fileobj.read()

        def download(self, name):
            return iter([stored[name]]) if name in stored else None

        def delete(self, name):
            stored.pop(name, None)

    monkeypatch.setattr(faults_router, "get_track_store", lambda: FakeStore())
    r = client.post(f"/faults/{fault_id}/photos", files={"file": ("a.jpg", JPEG_BYTES, "image/jpeg")})
    assert r.status_code == 201, r.text
    assert len(stored) == 1
    (name,) = stored
    assert name.startswith(f"fault-photos/{fault_id}/")
    assert list(local_store.rglob("*.jpg")) == []  # staging copy removed

    served = client.get(f"/faults/{fault_id}/photos/{r.json()['id']}")
    assert served.content == JPEG_BYTES

    assert client.delete(f"/faults/{fault_id}/photos/{r.json()['id']}").status_code == 204
    assert stored == {}


# ── work-order export ─────────────────────────────────────────────────────────

@pytest.fixture
def work_orders(db_session):
    ids = {
        "open": _seed_fault(db_session, panel_id="R1-C1", cls="hot-spot-high"),
        "assigned": _seed_fault(db_session, panel_id="R1-C2", cls="soiling", status="assigned",
                                assignee="Ravi"),
        "in_progress": _seed_fault(db_session, panel_id="R1-C3", cls="bypass-diode",
                                   status="in_progress", assignee="=HYPERLINK(\"x\")"),
        "resolved": _seed_fault(db_session, panel_id="R1-C4", cls="cell", status="resolved"),
        "stale": _seed_fault(db_session, panel_id="R1-C5", cls="module", status="stale"),
    }
    return ids


def _csv_rows(resp) -> list[dict]:
    return list(csv.DictReader(io.StringIO(resp.text)))


def test_work_orders_csv_defaults_to_actionable_statuses(client, work_orders):
    r = client.get("/parks/PARK_WF/work-orders?format=csv")
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("text/csv")
    assert "attachment" in r.headers["content-disposition"]
    rows = _csv_rows(r)
    assert {row["status"] for row in rows} == {"open", "assigned", "in_progress"}
    first = rows[0]
    for col in ("fault_id", "panel_id", "lat", "lon", "class", "severity", "priority",
                "status", "assignee", "due_date"):
        assert col in first
    hot = next(row for row in rows if row["class"] == "hot-spot-high")
    assert hot["lat"] == "18.5" and hot["lon"] == "73.8"
    assert hot["severity"] == SEVERITY_MAP["hot-spot-high"]
    assert hot["priority"] == "urgent"
    # most urgent first
    assert rows[-1]["class"] == "soiling"


def test_work_orders_csv_neutralises_formula_injection(client, work_orders):
    rows = _csv_rows(client.get("/parks/PARK_WF/work-orders?format=csv"))
    injected = next(row for row in rows if row["status"] == "in_progress")
    assert injected["assignee"].startswith("'=")


def test_work_orders_filter_by_status_and_assignee(client, work_orders):
    rows = _csv_rows(client.get("/parks/PARK_WF/work-orders?format=csv&status=assigned,in_progress"))
    assert {row["status"] for row in rows} == {"assigned", "in_progress"}

    rows = _csv_rows(client.get("/parks/PARK_WF/work-orders?format=csv&assignee=ravi"))
    assert [row["panel_id"] for row in rows] == ["R1-C2"]


def test_work_orders_rejects_bad_params(client, work_orders):
    assert client.get("/parks/PARK_WF/work-orders?format=pdf").status_code == 400
    assert client.get("/parks/PARK_WF/work-orders?status=bogus").status_code == 400
    assert client.get("/parks/NOPE/work-orders").status_code == 404


def test_work_orders_xlsx(client, work_orders):
    import openpyxl

    r = client.get("/parks/PARK_WF/work-orders?format=xlsx")
    assert r.status_code == 200, r.text
    assert "spreadsheetml" in r.headers["content-type"]
    ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
    header = [c.value for c in ws[1]]
    assert header[:4] == ["fault_id", "panel_id", "lat", "lon"]
    data = [dict(zip(header, [c.value for c in row])) for row in ws.iter_rows(min_row=2)]
    assert {d["status"] for d in data} == {"open", "assigned", "in_progress"}
    assert len(data) == 3


# ── migration ─────────────────────────────────────────────────────────────────

def _alembic_cfg(db_url: str):
    from pathlib import Path

    from alembic.config import Config

    repo_root = Path(__file__).resolve().parents[2]
    cfg = Config(str(repo_root / "alembic.ini"))
    cfg.set_main_option("script_location", str(repo_root / "alembic"))
    cfg.set_main_option("sqlalchemy.url", db_url)
    return cfg


def test_migration_0008_applies_on_sqlite_and_preserves_data(tmp_path, monkeypatch):
    import sqlalchemy as sa
    from alembic import command

    db_url = f"sqlite:///{tmp_path / 'mig.db'}"
    monkeypatch.setenv("AXALON_DB_URL", db_url)
    cfg = _alembic_cfg(db_url)

    command.upgrade(cfg, "0007")
    engine = sa.create_engine(db_url)
    with engine.begin() as conn:
        conn.execute(sa.text("INSERT INTO parks (id, name) VALUES ('P1', 'Park')"))
        conn.execute(sa.text(
            "INSERT INTO panel_faults (park_id, panel_id, class, status, notes) "
            "VALUES ('P1', 'R1-C1', 'soiling', 'stale', 'keep me')"
        ))

    command.upgrade(cfg, "head")
    insp = sa.inspect(engine)
    cols = {c["name"] for c in insp.get_columns("panel_faults")}
    assert {"assignee", "due_date", "priority", "resolved_at", "resolution_note"} <= cols
    assert insp.has_table("fault_photos")
    with engine.connect() as conn:
        row = conn.execute(sa.text("SELECT status, notes, assignee FROM panel_faults")).one()
    assert tuple(row) == ("stale", "keep me", None)

    command.downgrade(cfg, "0007")
    insp = sa.inspect(engine)
    assert "assignee" not in {c["name"] for c in insp.get_columns("panel_faults")}
    assert not insp.has_table("fault_photos")
    engine.dispose()


def test_startup_migrate_adds_workflow_columns_to_old_local_db(tmp_path):
    import sqlalchemy as sa

    from axalon.db.migrate import run_migrations

    engine = sa.create_engine(f"sqlite:///{tmp_path / 'old.db'}")
    with engine.begin() as conn:
        conn.execute(sa.text("CREATE TABLE parks (id VARCHAR PRIMARY KEY, name VARCHAR NOT NULL)"))
        conn.execute(sa.text(
            "CREATE TABLE panel_faults (id INTEGER PRIMARY KEY, park_id VARCHAR NOT NULL, "
            "panel_id VARCHAR NOT NULL, class VARCHAR NOT NULL, status VARCHAR)"
        ))

    actions = run_migrations(engine)

    cols = {c["name"] for c in sa.inspect(engine).get_columns("panel_faults")}
    assert {"assignee", "due_date", "priority", "resolved_at", "resolution_note"} <= cols
    assert "added panel_faults.assignee" in actions
    assert not any("panel_faults" in a for a in run_migrations(engine))  # idempotent
    engine.dispose()
